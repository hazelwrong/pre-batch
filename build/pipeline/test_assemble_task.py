"""The join between a finished orchestrator workspace and the builder's inputs.

Proves the step that was missing: agent outputs in, task data out, and a refusal
when anything upstream is stale.
"""
import json
import sys
import tempfile
import unittest
from pathlib import Path
from uuid import uuid4

import docx
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import assemble_task as ASM
import build_references as REF
import officestrip
from orchestrator import Pipeline, PipelineError, DEFAULT_POLICY

REFERENCE_SPEC = [
    {"filename": "PS-2026-04 Procurement Policy.md", "format": "md",
     "title": "Procurement Policy PS-2026-04",
     "blocks": [{"heading": "1. Scope"}, {"text": "Governs store systems."}]},
    {"filename": "Store Profile - Chaoyang Stores.xlsx", "format": "xlsx",
     "sheets": [{"name": "Stores", "columns": ["Store Code", "Store"],
                 "rows": [["CY-01", "Chaoyangmen"]]}]},
]
BLUEPRINT = {
    "sector": "Retail Trade",
    "occupation": "General and Operations Managers",
    "language": "en",
    "output_contract": ["Vendor Comparison.xlsx", "Recommendation to GM.docx"],
    "file_roles": {"policy": "PS-2026-04 Procurement Policy.md"},
}
DESIGN_NOTES = {
    "reasoning_points": ["an unquoted migration line"],
    "column_maps": {"*": {"A": 2}},
    "guards": {"domain_stopwords": ["store"]},
    "figure_pattern": r"RMB ([\d,]+\.\d{2})",
}
LINEAGE = {
    "statement": "Every figure originates in one of the reference files.",
    "input_population": {
        "PS-2026-04 Procurement Policy.md": {"role": "authoritative rules"},
        "Store Profile - Chaoyang Stores.xlsx": {"role": "estate register"},
    },
}


def rubric_items(n=40):
    items = [{"score": 2, "criterion": "Item %d states one checkable thing" % i,
              "rubric_item_id": str(uuid4()), "required": True,
              "verification": "Read it."} for i in range(30)]
    items += [{"score": 4, "criterion": "Judgement item %d" % i,
               "rubric_item_id": str(uuid4()), "required": True,
               "verification": "Judge it."} for i in range(10)]
    return items


class AssembleTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name)
        self.tasks = self.base / "tasks"
        self.staging = self.base / "staging"
        self.pipeline = Pipeline.initialise(self.base / "work", str(uuid4()))
        self.intake = self.base / "intake"
        self.intake.mkdir()

        # references are built by S-REF from the same spec the designer writes
        refs = self.intake / "references"
        REF.build(REFERENCE_SPEC, refs)
        for category in ("occupation_standard", "material_pool", "source_manifest",
                         "coverage"):
            folder = self.intake / category
            folder.mkdir()
            (folder / (category + ".txt")).write_text(category, encoding="utf-8")
            self.pipeline.add_artifact(category, [folder])
        self.pipeline.add_artifact("references", [refs])
        self.pipeline.add_artifact("policy", [DEFAULT_POLICY])

    def tearDown(self):
        self.tmp.cleanup()

    def complete(self, role, agent, context, outputs):
        run_root = self.pipeline.prepare(role, agent, context)
        for category, value in outputs.items():
            folder = run_root / "output" / category
            folder.mkdir()
            if category == "prompt":
                (folder / "prompt.md").write_text(value, encoding="utf-8")
            elif category == "rubric":
                (folder / "rubric.json").write_text(json.dumps(value), encoding="utf-8")
            elif category == "reference_spec":
                (folder / "spec.json").write_text(json.dumps(value), encoding="utf-8")
            elif category == "gold":
                workbook = folder / BLUEPRINT["output_contract"][0]
                book = openpyxl.Workbook()
                book.active["A1"] = "gold"
                book.save(workbook)
                officestrip.strip_and_pin(str(workbook))
                document = folder / BLUEPRINT["output_contract"][1]
                doc = docx.Document()
                doc.add_paragraph("gold")
                doc.save(document)
                officestrip.strip_and_pin(str(document))
            elif isinstance(value, (dict, list)):
                (folder / "report.json").write_text(json.dumps(value), encoding="utf-8")
            else:
                (folder / "artifact.txt").write_text(value, encoding="utf-8")
        run_id = json.loads((run_root / "run_contract.json").read_text())["run_id"]
        self.pipeline.submit(run_id, "passed")

    def build_all(self, lineage=None):
        self.complete("gold_curator", "curator", "ctx-gold", {
            "gold": None,
            "gold_provenance": {
                "source_type": "desensitization",
                "production_method": "supplier work record, de-identified",
                "is_real_deliverable": True,
                "real_deliverable_files": [{
                    "filename": BLUEPRINT["output_contract"][0],
                    "source_url": "https://example.test/gold.xlsx",
                    "source_sha256": "a" * 64,
                }],
                "rights_holder": "Supplier Ltd",
                "license": "Supplier grants evaluation and redistribution rights",
                "usage_scope": "GDPval evaluation and redistribution"},
            "production_notes": "notes"})
        self.complete("task_designer", "designer", "ctx-design", {
            "task_blueprint": BLUEPRINT, "design_notes": DESIGN_NOTES,
            "reference_spec": REFERENCE_SPEC,
            "lineage_draft": lineage or LINEAGE})
        self.complete("prompt_author", "author", "ctx-prompt", {
            "prompt": "You are the operations manager. Produce the workbook.",
            "output_contract": {"files": BLUEPRINT["output_contract"]}})
        self.complete("solver", "solver", "ctx-solver", {
            "solver_deliverables": "solution",
            "solver_report": {"prompt_self_contained": True, "solvable": True,
                              "task_multistep": True, "separating_power": "sufficient",
                              "difficulty_evidence": ["cross-file reconciliation"],
                              "blocking_ambiguities": []}})
        self.complete("verifier", "verifier", "ctx-verifier", {
            "expected_values": {"values": [{"name": "stores", "value": 3}]},
            "verifier_report": {"recompute_passed": True, "lineage_valid": True,
                                "mismatches": [],
                                "demands_without_landing_place": []}})
        self.complete("rubric", "rubric", "ctx-rubric", {"rubric": rubric_items()})

    def run_assemble(self):
        return ASM.assemble(self.pipeline.root, self.tasks, self.staging)

    def test_assembles_task_data_and_staging(self):
        self.build_all()
        result = self.run_assemble()
        out = Path(result["task_dir"])
        for name in ("prompt.md", "rubric.json", "rubric_pretty.txt",
                     "expected_values.json", "lineage.json", "task_meta.json",
                     "reference_spec.json", "gold_provenance.json"):
            self.assertTrue((out / name).is_file(), name)

        meta = json.loads((out / "task_meta.json").read_text(encoding="utf-8"))
        self.assertEqual(meta["sector"], "Retail Trade")
        # The evaluator-only half reaches task_meta, not the prompt author.
        self.assertEqual(meta["column_maps"], DESIGN_NOTES["column_maps"])
        self.assertEqual(meta["figure_pattern"], DESIGN_NOTES["figure_pattern"])
        self.assertEqual(meta["guards"], DESIGN_NOTES["guards"])
        # File order follows the spec, not the filesystem's sort.
        self.assertEqual(meta["file_order"]["reference_files"],
                         [e["filename"] for e in REFERENCE_SPEC])
        self.assertEqual(meta["file_order"]["deliverable_files"],
                         BLUEPRINT["output_contract"])

        staged = sorted(p.name for p in
                        (self.staging / result["task_id"] / "reference_files").iterdir())
        self.assertEqual(staged, sorted(e["filename"] for e in REFERENCE_SPEC))
        self.assertEqual(
            sorted(p.name for p in
                   (self.staging / result["task_id"] / "deliverable_files").iterdir()),
            sorted(BLUEPRINT["output_contract"]))

        # The marking sheet is a person's work and is reported as outstanding
        # rather than invented.
        self.assertEqual(result["awaiting_human"], ["gold_marking.json"])

    def test_rubric_pretty_is_derived_from_the_items(self):
        self.build_all()
        out = Path(self.run_assemble()["task_dir"])
        items = json.loads((out / "rubric.json").read_text(encoding="utf-8"))
        pretty = (out / "rubric_pretty.txt").read_text(encoding="utf-8").rstrip("\n")
        self.assertEqual(pretty.split("\n\n")[0],
                         "[+%d] %s" % (items[0]["score"], items[0]["criterion"]))
        self.assertEqual(len(pretty.split("\n\n")), len(items))

    def test_expected_values_accepts_descriptive_filename(self):
        folder = self.intake / "expected_values_named"
        folder.mkdir()
        expected = {"values": [{"name": "events", "value": 15}]}
        (folder / "expected_values.json").write_text(
            json.dumps(expected), encoding="utf-8")
        self.pipeline.add_artifact("expected_values", [folder])

        self.assertEqual(
            ASM._read(self.pipeline, "expected_values", "report.json"), expected)

    def test_lineage_accepts_structured_input_universe(self):
        references = [entry["filename"] for entry in REFERENCE_SPEC]
        lineage = {
            "input_universe": [
                {"file": name, "role": "source", "join_key": "record_id"}
                for name in references
            ]
        }

        ok, detail = ASM.LN.verify(lineage, references)

        self.assertTrue(ok, detail)

    def test_lineage_accepts_reference_reconstruction_entries(self):
        references = [entry["filename"] for entry in REFERENCE_SPEC]
        lineage = {
            "reference_reconstruction_lineage": [
                {"reference": name, "gold_origin": "observable output region"}
                for name in references
            ]
        }

        ok, detail = ASM.LN.verify(lineage, references)

        self.assertTrue(ok, detail)

    def test_guard_rule_list_is_preserved_in_task_meta_shape(self):
        rules = [{"id": "G-01", "rule": "Do not infer missing evidence."}]

        self.assertEqual(ASM._normalise_guards(rules), {"rules": rules})

    def test_file_role_list_is_preserved_without_triggering_template_roles(self):
        roles = [{"file": "Events.csv", "role": "acceptance history"}]

        self.assertEqual(
            ASM._normalise_file_roles(roles), {"source_files": roles})

    def test_stale_upstream_refuses_to_assemble(self):
        self.build_all()
        replacement = self.base / "revised-prompt.md"
        replacement.write_text("revised", encoding="utf-8")
        self.pipeline.add_artifact("prompt", [replacement])
        with self.assertRaises(PipelineError) as caught:
            self.run_assemble()
        self.assertIn("not current", str(caught.exception))

    def test_lineage_must_cover_the_staged_references(self):
        # Built through with the short lineage from the start. Patching it in
        # afterwards would only prove the staleness rule, which is a different
        # test: once the designer re-runs, everything downstream is stale too.
        short = dict(LINEAGE, input_population={
            "PS-2026-04 Procurement Policy.md": {"role": "authoritative rules"}})
        self.build_all(lineage=short)
        with self.assertRaises(PipelineError) as caught:
            self.run_assemble()
        self.assertIn("untraced", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
